import os
import re
import cv2
import io
import numpy as np
import streamlit as st
from PIL import Image
import pytesseract
import pypdf
import docx2txt
from docx import Document  
from fpdf import FPDF  
from deep_translator import GoogleTranslator  
from pdf2image import convert_from_bytes  

# 1. Configuration Tesseract
chemin_tesseract = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
pytesseract.pytesseract.tesseract_cmd = chemin_tesseract
os.environ['TESSDATA_PREFIX'] = r'C:\Program Files\Tesseract-OCR\tessdata'

# Configurer la page
st.set_page_config(page_title="OCR de Babacar", layout="wide", initial_sidebar_state="collapsed")

# --- INITIALISATION DE LA MÉMOIRE (Session State) ---
if "texte_extrait" not in st.session_state:
    st.session_state.texte_extrait = ""
if "fichier_actuel" not in st.session_state:
    st.session_state.fichier_actuel = None

# --- DESIGN & CSS STYLE SCRIPT ---
st.markdown("""
    <style>
    #MainMenu {visibility: hidden;}
    footer {visibility: hidden;}
    header {visibility: hidden;}
    
    div[data-testid="stFileUploader"] section div div {
        display: none !important;
    }
    div[data-testid="stFileUploaderDropzone"] small {
        display: none !important;
    }
    
    [data-testid="sidebar-toggle"] {
        display: none !important;
    }
    section[data-testid="stSidebar"] {
        display: none !important;
    }
    
    .stApp {
        background: linear-gradient(135deg, #0f0c20 0%, #15102a 100%);
        color: #ffffff;
    }
    
    .main-title {
        font-family: 'Helvetica Neue', sans-serif;
        font-weight: 800;
        font-size: 3.2rem;
        background: linear-gradient(45deg, #ff007f, #7f00ff);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        text-align: center;
        margin-top: 20px;
        margin-bottom: 5px;
    }
    
    .subtitle {
        text-align: center;
        color: #b3b0cb;
        font-size: 1.2rem;
        margin-bottom: 50px;
    }

    div[data-testid="stVerticalBlock"] > div:has(div.stImage) {
        background-color: #1e1938;
        border-radius: 15px;
        padding: 25px;
        box-shadow: 0 8px 32px 0 rgba(0, 0, 0, 0.3);
        border: 1px solid rgba(255, 255, 255, 0.05);
    }
    
    .stFileUploader {
        background-color: #1e1938;
        border: 2px dashed #7f00ff;
        border-radius: 12px;
        padding: 20px;
        max-width: 800px;
        margin: 0 auto 30px auto;
    }

    .stButton > button {
        background: linear-gradient(45deg, #7f00ff, #ff007f) !important;
        color: white !important;
        font-weight: bold !important;
        border: none !important;
        border-radius: 8px !important;
        padding: 14px 28px !important;
        box-shadow: 0 4px 15px rgba(127, 0, 255, 0.4) !important;
        transition: all 0.3s ease !important;
        width: 100%;
        font-size: 1.1rem !important;
    }
    
    .stButton > button:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 6px 20px rgba(255, 0, 127, 0.6) !important;
    }
    
    div.stDownloadButton > button {
        background: #1e1938 !important;
        color: #00ffcc !important;
        border: 1px solid #3d356b !important;
        box-shadow: none !important;
        font-size: 0.95rem !important;
        padding: 10px 15px !important;
        margin-bottom: 10px;
    }
    div.stDownloadButton > button:hover {
        background: #3d356b !important;
        border-color: #00ffcc !important;
    }
    
    .stTextArea textarea {
        background-color: #0d0a1b !important;
        color: #00ffcc !important;
        font-family: 'Courier New', monospace !important;
        border: 1px solid #3d356b !important;
        border-radius: 10px !important;
        font-size: 1.05rem !important;
    }

    /* Style de la zone de prévisualisation avec surlignage HTML */
    .preview-box {
        background-color: #0d0a1b;
        color: #ffffff;
        font-family: 'Courier New', monospace;
        border: 1px solid #3d356b;
        border-radius: 10px;
        padding: 15px;
        max-height: 250px;
        overflow-y: auto;
        white-space: pre-wrap;
        margin-top: 10px;
        font-size: 1.05rem;
    }
    .highlight {
        background-color: #ffcc00 !important;
        color: #000000 !important;
        font-weight: bold;
        border-radius: 3px;
        padding: 0 2px;
    }
    </style>
""", unsafe_allow_html=True)

st.markdown('<h1 class="main-title">OCR de Babacar</h1>', unsafe_allow_html=True)
st.markdown('<p class="subtitle">Plateforme professionnelle de numérisation, traduction et analyse de documents</p>', unsafe_allow_html=True)

st.markdown("<div style='text-align: center; font-weight: bold; margin-bottom: 10px;'>📂 Déposez votre document (Image, PDF natif/scanné, Word...)</div>", unsafe_allow_html=True)
fichier_uploade = st.file_uploader("", type=["png", "jpg", "jpeg", "jfif", "webp", "pdf", "docx"])

if fichier_uploade is not None and fichier_uploade.name != st.session_state.fichier_actuel:
    st.session_state.texte_extrait = ""
    st.session_state.fichier_actuel = fichier_uploade.name

if fichier_uploade is not None:
    col1, col2 = st.columns(2, gap="large")
    
    nom_base, extension = os.path.splitext(fichier_uploade.name)
    extension = extension.lower()

    with col1:
        st.markdown("<h4 style='color:#7f00ff; margin-bottom: 15px;'>🖼️ Aperçu du Document</h4>", unsafe_allow_html=True)
        if extension in [".pdf", ".docx"]:
            st.info(f"📄 Fichier {extension.upper()} détecté : **{fichier_uploade.name}**")
            st.write("Prêt pour l'extraction de texte.")
        else:
            image_pil = Image.open(fichier_uploade)
            st.image(image_pil, use_container_width=True)

    with col2:
        st.markdown("<h4 style='color:#ff007f; margin-bottom: 15px;'>⚡ Analyse</h4>", unsafe_allow_html=True)
        
        if st.button("🚀 EXTRAIRE LES DONNÉES"):
            with st.spinner("Analyse et lecture du document en cours..."):
                texte_brut = ""
                
                if extension == ".docx":
                    texte_brut = docx2txt.process(fichier_uploade)
                
                elif extension == ".pdf":
                    octets_pdf = fichier_uploade.read()
                    lecteur_pdf = pypdf.PdfReader(io.BytesIO(octets_pdf))
                    pages_texte = []
                    
                    for page in lecteur_pdf.pages:
                        texte_page = page.extract_text()
                        if texte_page:
                            pages_texte.append(texte_page)
                    texte_brut = "\n".join(pages_texte)
                    
                    if not texte_brut.strip():
                        st.info("ℹ️ PDF scanné détecté. Conversion des pages en images...")
                        try:
                            images_pdf = convert_from_bytes(octets_pdf)
                            textes_ocr_pages = []
                            for img in images_pdf:
                                img_cv = np.array(img)
                                img_gris = cv2.cvtColor(img_cv, cv2.COLOR_RGB2GRAY)
                                kernel = np.array([[-1,-1,-1], [-1,9,-1], [-1,-1,-1]])
                                img_gris = cv2.filter2D(img_gris, -1, kernel)
                                img_traitee = cv2.threshold(img_gris, 0, 255, cv2.THRESH_BINARY | cv2.THRESH_OTSU)[1]
                                page_txt = pytesseract.image_to_string(Image.fromarray(img_traitee), lang='eng+fra', config='--psm 6')
                                textes_ocr_pages.append(page_txt)
                            texte_brut = "\n".join(textes_ocr_pages)
                        except Exception:
                            texte_brut = "ERREUR_PDF_IMAGE"
                else:
                    img_cv = np.array(image_pil)
                    if len(img_cv.shape) == 3:
                        if img_cv.shape[2] == 4:
                            img_cv = cv2.cvtColor(img_cv, cv2.COLOR_RGBA2BGR)
                        elif img_cv.shape[2] == 3:
                            img_cv = cv2.cvtColor(img_cv, cv2.COLOR_RGB2BGR)

                    if len(img_cv.shape) == 3 and img_cv.shape[2] == 3:
                        img_gris = cv2.cvtColor(img_cv, cv2.COLOR_BGR2GRAY)
                    else:
                        img_gris = img_cv.copy()
                        
                    kernel = np.array([[-1,-1,-1], [-1,9,-1], [-1,-1,-1]])
                    img_gris = cv2.filter2D(img_gris, -1, kernel)
                    img_traitee = cv2.threshold(img_gris, 0, 255, cv2.THRESH_BINARY | cv2.THRESH_OTSU)[1]
                    image_a_scanner = Image.fromarray(img_traitee)
                    texte_brut = pytesseract.image_to_string(image_a_scanner, lang='eng+fra', config='--psm 6')

                if texte_brut not in ["ERREUR_PDF_IMAGE"]:
                    lignes_nettoyees = []
                    for ligne in texte_brut.split('\n'):
                        ligne_propre = re.sub(r'^[=\s«»•\-\*]+', '', ligne).strip()
                        lignes_nettoyees.append(ligne_propre)
                    st.session_state.texte_extrait = '\n'.join(lignes_nettoyees)
                else:
                    st.session_state.texte_extrait = "ERREUR"

        # --- ZONE D'AFFICHAGE ET OPTIONS INTÉGRÉES ---
        if st.session_state.texte_extrait.strip() and st.session_state.texte_extrait != "ERREUR":
            st.success("🎉 Analyse complétée avec succès !")
            
            # 🔍 --- BLOC DE RECHERCHE DYNAMIQUE ---
            st.markdown("<h5 style='color:#00ffcc; margin-top:10px;'>🔍 Recherche de mots-clés</h5>", unsafe_allow_html=True)
            mot_recherche = st.text_input("Tapez un mot ou une phrase à chercher :", key="search_input")
            
            # Zone d'édition interactive
            st.session_state.texte_extrait = st.text_area(
                "📝 Modifier ou corriger le texte extrait :", 
                value=st.session_state.texte_extrait, 
                height=200
            )

            # Si une recherche est en cours, on calcule et on montre le surlignage HTML juste en dessous
            if mot_recherche.strip():
                # Compter le nombre d'occurrences (insensible à la casse)
                occurrences = len(re.findall(re.escape(mot_recherche), st.session_state.texte_extrait, re.IGNORECASE))
                
                if occurrences > 0:
                    st.markdown(f"📊 **{occurrences}** occurrence(s) trouvée(s) pour le mot : `{mot_recherche}`")
                    # Surlignage dynamique via regex HTML
                    pattern = re.compile(rf"({re.escape(mot_recherche)})", re.IGNORECASE)
                    texte_surligne = pattern.sub(r'<span class="highlight">\1</span>', st.session_state.texte_extrait)
                    # Affichage du résultat surligné
                    st.markdown(f'<div class="preview-box">{texte_surligne}</div>', unsafe_allow_html=True)
                else:
                    st.warning(f"Aucun résultat trouvé pour `{mot_recherche}`.")
            
            # --- BLOC DE TRADUCTION AUTOMATIQUE ---
            st.markdown("<h5 style='color:#ff007f; margin-top:20px;'>🌍 Traduction Automatique</h5>", unsafe_allow_html=True)
            langues_dispo = {"Anglais": "en", "Français": "fr", "Espagnol": "es", "Arabe": "ar", "Allemand": "de"}
            
            lang_col1, lang_col2 = st.columns([2, 1])
            with lang_col1:
                langue_cible = st.selectbox("Choisir la langue de destination :", list(langues_dispo.keys()))
            with lang_col2:
                st.write("")
                st.write("")
                if st.button("🔄 Traduire"):
                    with st.spinner("Traduction..."):
                        try:
                            code_langue = langues_dispo[langue_cible]
                            texte_traduit = GoogleTranslator(source='auto', target=code_langue).translate(st.session_state.texte_extrait)
                            st.session_state.texte_extrait = texte_traduit
                            st.rerun()
                        except Exception:
                            st.error("Erreur réseau lors de la traduction.")

            # --- ZONE D'EXPORTATION ---
            st.markdown("<h5 style='color:#00ffcc; margin-top:25px;'>📥 Options d'exportation :</h5>", unsafe_allow_html=True)
            exp_col1, exp_col2, exp_col3 = st.columns(3)
            
            with exp_col1:
                st.download_button(
                    label="📄 Fichier Texte (.txt)", data=st.session_state.texte_extrait,
                    file_name=f"{nom_base}.txt", mime="text/plain", use_container_width=True
                )
            
            with exp_col2:
                doc = Document()
                doc.add_paragraph(st.session_state.texte_extrait)
                buffer_word = io.BytesIO()
                doc.save(buffer_word)
                buffer_word.seek(0)
                st.download_button(
                    label="📝 Document Word (.docx)", data=buffer_word,
                    file_name=f"{nom_base}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True
                )
            
            with exp_col3:
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Arial", size=11)
                texte_pdf = st.session_state.texte_extrait.encode('latin-1', 'replace').decode('latin-1')
                for ligne in texte_pdf.split('\n'):
                    ligne_nettoyee = ligne.strip()
                    if not ligne_nettoyee:
                        pdf.ln(6) 
                    else:
                        try:
                            pdf.multi_cell(w=0, h=6, txt=ligne_nettoyee)
                        except Exception:
                            pdf.cell(w=0, h=6, txt=ligne_nettoyee[:50] + "...")
                            pdf.ln(6)
                buffer_pdf = io.BytesIO()
                pdf.output(buffer_pdf)
                buffer_pdf.seek(0)
                st.download_button(
                    label="📕 Document PDF (.pdf)", data=buffer_pdf,
                    file_name=f"{nom_base}.pdf", mime="application/pdf", use_container_width=True
                )

        elif st.session_state.texte_extrait == "ERREUR":
            st.error("❌ Échec de l'extraction. Impossible d'analyser ce document.")